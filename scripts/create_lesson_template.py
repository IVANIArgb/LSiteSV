#!/usr/bin/env python3
"""Создать lesson_template.docx с инструкцией и примерами тегов."""

from __future__ import annotations

import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
OUT = PROJECT_ROOT / "docs" / "lesson_template.docx"


def main() -> int:
    try:
        from docx import Document
    except ImportError:
        print("pip install python-docx")
        return 1

    doc = Document()
    doc.add_heading("Шаблон урока LMS", level=1)
    doc.add_paragraph(
        "Сохраните этот файл как lesson.docx в папке урока и запустите "
        "пункт 10 в lms_manager.bat или: python scripts/parse_word_lesson.py <папка_урока>"
    )
    doc.add_heading("Текст и заголовки", level=2)
    doc.add_paragraph("# Заголовок урока (строка с # станет блоком heading)")
    doc.add_paragraph("Обычный текст урока.")
    doc.add_heading("Медиа-теги", level=2)
    doc.add_paragraph("[IMAGE: schema.png] — положите schema.png в папку images/")
    doc.add_paragraph("[VIDEO: intro.mp4] — положите intro.mp4 в videos/")
    doc.add_paragraph("[FILE: manual.pdf] — положите manual.pdf в files/")
    doc.add_heading("Тест", level=2)
    doc.add_paragraph("=== ТЕСТ НАЧАЛО ===")
    doc.add_paragraph("TITLE: Проверочный тест")
    doc.add_paragraph("Q: Столица России?")
    doc.add_paragraph("A) Москва")
    doc.add_paragraph("B) Париж")
    doc.add_paragraph("TYPE: single")
    doc.add_paragraph("CORRECT: A")
    doc.add_paragraph("POINTS: 1")
    doc.add_paragraph("=== ТЕСТ КОНЕЦ ===")
    doc.add_paragraph("Текст после теста.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Создан: {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
